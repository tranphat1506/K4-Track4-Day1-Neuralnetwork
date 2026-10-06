"""results_table.py — PSEUDO-CODE."""
from __future__ import annotations

import json
from pathlib import Path


def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result ra file json."""
    path = Path(results_dir)
    path.mkdir(parents=True, exist_ok=True)
    exp_id = result["cfg"]["exp_id"]
    file_path = path / f"{exp_id}.json"
    
    data_to_save = {
        "cfg": result["cfg"],
        "history": result["history"],
        "summary": result["summary"]
    }
    
    with open(file_path, "w", encoding="utf-8") as f:
        json.dump(data_to_save, f, indent=2)
        
    return str(file_path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json."""
    path = Path(results_dir)
    if not path.exists():
        return []
        
    results = []
    for file_path in sorted(path.glob("*.json")):
        with open(file_path, "r", encoding="utf-8") as f:
            results.append(json.load(f))
    return results


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng."""
    row = {}
    row.update(result["cfg"])
    row["hidden"] = str(row["hidden"])  # Excel có thể không hỗ trợ list/tuple trực tiếp tốt
    row.update(result["summary"])
    
    if eval_scores:
        row["eval_acc"] = eval_scores.get("acc")
        row["eval_macro_f1"] = eval_scores.get("macro_f1")
    else:
        row["eval_acc"] = None
        row["eval_macro_f1"] = None
        
    row["figure_file"] = f"figures/{result['cfg']['exp_id']}.png"
    row["notes"] = notes
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu."""
    import openpyxl
    wb = openpyxl.load_workbook(template_path)
    ws = wb["Experiments"]
    
    headers = [cell.value for cell in ws[1]]
    
    start_row = 2
    for i, row_dict in enumerate(rows):
        row_idx = start_row + i
        for col_idx, header in enumerate(headers, start=1):
            if header in row_dict and row_dict[header] is not None:
                ws.cell(row=row_idx, column=col_idx, value=row_dict[header])
                
    wb.save(out_path)
